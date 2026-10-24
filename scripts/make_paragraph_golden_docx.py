"""Write contracts/goldens/paragraphs/synthetic/source.docx: one deliberate case per DOCX_EXTRACTION rule.

    PYTHONPATH=src .venv/bin/python scripts/make_paragraph_golden_docx.py

A one-off generator (the committed .docx is the golden; this only records how it was made). Every line
of the body below names the rule it exercises, so a reader of extracted.txt can map each line back.
"""
import io
import os
import zipfile

import docx
from lxml import etree

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "contracts", "goldens", "paragraphs", "synthetic", "source.docx")

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = (f'xmlns:w="{W}" xmlns:v="urn:schemas-microsoft-com:vml" '
      'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')


def p(inner: str) -> str:
    return f"<w:p>{inner}</w:p>"


def r(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def tc(inner: str, props: str = "") -> str:
    return f"<w:tc><w:tcPr>{props}</w:tcPr>{inner}</w:tc>"


BODY = "".join([
    p(r("1. A plain paragraph.")),
    p("<w:pPr><w:tabs><w:tab w:val='left' w:pos='720'/></w:tabs></w:pPr>"
      + r("2. Tab stop definitions") + "<w:r><w:tab/></w:r>" + r("are not text; a run tab is one space.")),
    p(r("3. A line break") + "<w:r><w:br/></w:r>" + r("inside a paragraph is one space.")),
    p(r("   ")),                                                            # whitespace-only: skipped
    p(r("4. Runs   with    spaced   words collapse.")),
    p(r("5. Inserted text ") + "<w:ins w:id='1' w:author='a'>" + r("is kept") + "</w:ins>"
      + "<w:del w:id='2' w:author='a'><w:r><w:delText>DELETED</w:delText></w:r></w:del>"
      + "<w:moveFrom w:id='3' w:author='a'>" + r("MOVEDAWAY") + "</w:moveFrom>"
      + "<w:moveTo w:id='4' w:author='a'>" + r(" and moved-to text too.") + "</w:moveTo>"),
    p(r("6. A ") + "<w:hyperlink r:id='rIdX'>" + r("hyperlink's text") + "</w:hyperlink>" + r(" is kept.")),
    p(r("7. A simple field result: ") + "<w:fldSimple w:instr=' PAGE '>" + r("42") + "</w:fldSimple>"),
    p(r("8. A complex field: ") + "<w:r><w:fldChar w:fldCharType='begin'/></w:r>"
      + "<w:r><w:instrText xml:space='preserve'> DATE \\@ yyyy </w:instrText></w:r>"
      + "<w:r><w:fldChar w:fldCharType='separate'/></w:r>" + r("2026")
      + "<w:r><w:fldChar w:fldCharType='end'/></w:r>"),
    p(r("9. Inline content control: ") + "<w:sdt><w:sdtContent>" + r("kept") + "</w:sdtContent></w:sdt>"),
    "<w:sdt><w:sdtContent>" + p(r("10. A body-level content control is walked in place.")) + "</w:sdtContent></w:sdt>",
    p(r("11. A text box is outside the flow")
      + "<w:r><w:pict><v:shape><v:textbox><w:txbxContent>" + p(r("TEXTBOX")) + "</w:txbxContent>"
      + "</v:textbox></v:shape></w:pict></w:r>"
      + "<mc:AlternateContent><mc:Choice Requires='wps'><w:r><w:t>CHOICE</w:t></w:r></mc:Choice>"
      + "<mc:Fallback><w:r><w:t>FALLBACK</w:t></w:r></mc:Fallback></mc:AlternateContent>"),
    "<w:tbl><w:tblGrid><w:gridCol/><w:gridCol/><w:gridCol/></w:tblGrid>",
    "<w:tr>" + tc(p(r("12a. header"))) + tc(p(r("12b. spans two")), "<w:gridSpan w:val='2'/>") + "</w:tr>",
    "<w:tr>" + tc(p(r("13a. merge start")), "<w:vMerge w:val='restart'/>") + tc(p(r(""))) + tc(p(r("13c."))) + "</w:tr>",
    "<w:tr>" + tc(p(r("CONTINUATION")), "<w:vMerge/>") + tc(p(r("14b. two")) + p(r("paragraphs"))) + tc(
        p(r("14c. outer")) + "<w:tbl><w:tr>" + tc(p(r("NESTED"))) + "</w:tr></w:tbl>") + "</w:tr>",
    "<w:tr>" + tc(p(r(""))) + tc(p(r(" "))) + tc(p(r(""))) + "</w:tr>",       # all empty: skipped
    "<w:tr><w:sdt><w:sdtContent>" + tc(p(r("15a. cell in a control"))) + "</w:sdtContent></w:sdt>"
    + tc(p(r("15b."))) + tc(p(r("15c."))) + "</w:tr>",
    "</w:tbl>",
    p(r("16. After the table, in document order.")),
])


def main() -> None:
    d = docx.Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "HEADER"
    sec.footer.paragraphs[0].text = "FOOTER"
    buf = io.BytesIO()
    d.save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    doc = etree.fromstring(zin.read("word/document.xml"))
    body = doc.find(f"{{{W}}}body")
    sect = body.find(f"{{{W}}}sectPr")
    for child in list(body):
        if child is not sect:
            body.remove(child)
    frag = etree.fromstring(f"<w:body {NS}>{BODY}</w:body>")
    for i, child in enumerate(frag):
        body.insert(i, child)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = etree.tostring(doc, xml_declaration=True, encoding="UTF-8", standalone=True) \
                if item.filename == "word/document.xml" else zin.read(item.filename)
            zout.writestr(item, data)
    print(f"wrote {os.path.relpath(OUT, ROOT)}")


if __name__ == "__main__":
    main()
