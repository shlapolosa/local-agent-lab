"""The workbook the team edits, and the master it turns back into.

The team asked for a spreadsheet and that is what they get — but the workbook is an EXCHANGE
format, never the source. DR-02 requires the human-readable master and the agent-readable form to
be the same artifact at the same version (`derived_from = master_sha256`), and that link is only
true if the derivation actually happens. So the round trip is: master -> workbook -> master, and
what the publisher hashes is the REGENERATED master. A workbook that cannot be turned back into
the master it came from is refused here rather than discovered at publish.
"""
import importlib.util
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[3]
spec = importlib.util.spec_from_file_location("artifacts_workbook",
                                              ROOT / "scripts" / "artifacts_workbook.py")
wb = importlib.util.module_from_spec(spec)
spec.loader.exec_module(wb)

from lab.core.reference import master  # noqa: E402

SAMPLE = """# Criticality taxonomy

**Artifact:** criticality_taxonomy
**Source:** CAFE_Artifacts_Visualisation_v0_25.html

| Class | Dominant failure | Rigor it sets |
|---|---|---|
| Routine | Someone is inconvenienced | Standard harness |
| Business-critical | Financial loss | Golden dataset; named owner |
"""


def test_a_sheet_carries_the_schema_a_reader_needs_before_the_values():
    """A column of values with no schema is a column somebody guesses at. The record type, the key
    fields and the owner are what say which column identifies a row and who may change it."""
    rows = wb.sheet_rows("criticality-taxonomy", master.parse(SAMPLE),
                         ("class", "Class", "risk"))
    flat = [str(c) for row in rows for c in row]
    assert "criticality-taxonomy" in flat and "class" in flat and "risk" in flat
    assert "Class" in flat, "the key field must be named, not merely be a column"
    assert "Routine" in flat and "Business-critical" in flat


def test_the_table_header_row_is_the_master_s_own_headers():
    rows = wb.sheet_rows("criticality-taxonomy", master.parse(SAMPLE), ("class", "Class", "risk"))
    assert ["Class", "Dominant failure", "Rigor it sets"] in [list(map(str, r)) for r in rows]


def test_a_workbook_round_trips_to_the_master_it_came_from(tmp_path):
    """The whole justification for accepting a spreadsheet: what publishes is a REGENERATED master,
    byte-identical in content to a master parsed from the workbook. If this ever fails, the
    workbook has become a second source of truth and DR-02 no longer holds."""
    parsed = master.parse(SAMPLE)
    path = tmp_path / "artifacts.xlsx"
    wb.export({"criticality-taxonomy": (parsed, ("class", "Class", "risk"))}, path)
    back = wb.read_sheet(path, "criticality-taxonomy")
    assert back.headers == parsed.headers
    assert back.rows == parsed.rows
    assert back.meta.get("Artifact") == parsed.meta.get("Artifact")


def test_a_sheet_whose_header_row_was_edited_is_refused(tmp_path):
    """Renaming a column silently changes what `key_fields` points at, and a record keyed on a
    column that no longer exists is one the publisher cannot match to its predecessor."""
    parsed = master.parse(SAMPLE)
    path = tmp_path / "a.xlsx"
    wb.export({"criticality-taxonomy": (parsed, ("class", "Class", "risk"))}, path)
    import openpyxl
    book = openpyxl.load_workbook(path)
    sheet = book["criticality-taxonomy"]
    header_at = wb.header_row_index(sheet)
    sheet.cell(row=header_at, column=1, value="Klass")
    book.save(path)
    with pytest.raises(wb.WorkbookError, match="key field"):
        wb.read_sheet(path, "criticality-taxonomy", key_fields=("Class",))


def test_an_index_sheet_lists_every_artifact_so_the_team_can_see_the_whole_ask(tmp_path):
    path = tmp_path / "a.xlsx"
    wb.export({"criticality-taxonomy": (master.parse(SAMPLE), ("class", "Class", "risk"))}, path)
    import openpyxl
    book = openpyxl.load_workbook(path)
    assert wb.INDEX in book.sheetnames
    flat = [str(c.value) for row in book[wb.INDEX].iter_rows() for c in row]
    assert "criticality-taxonomy" in flat and "risk" in flat


COLLIDING = """# Input artifacts

**Artifact:** input_artifacts

| Artifact | Owner | Consumed at |
|---|---|---|
| Business capability map | Enterprise architecture | Match capabilities |
"""


def test_a_table_whose_first_column_is_named_like_a_schema_label_still_round_trips(tmp_path):
    """Four of the 53 real artifacts head their table with `Artifact` or `Owner` — the same words
    the schema block uses as labels. Locating the header row by "first row that is not a label"
    therefore skipped it and read the first DATA row as headers, and the import refused them.

    The separator is what is reliable: `export` writes exactly one blank row between the schema
    block and the table, so the header is the first non-empty row after it. No guessing at names.
    """
    parsed = master.parse(COLLIDING)
    path = tmp_path / "a.xlsx"
    wb.export({"input-artifacts": (parsed, ("artifact", "Artifact", "EA"))}, path)
    back = wb.read_sheet(path, "input-artifacts", key_fields=("Artifact",))
    assert back.headers == ("Artifact", "Owner", "Consumed at")
    assert back.rows == parsed.rows


UNNAMED_FIRST = """# Risk derivation

**Artifact:** risk_derivation
**Source:** CAFE_Artifacts_Visualisation_v0_25.html
**Section:** properties

|  | Exposure | Influence |
|---|---|---|
| Question | How bad is it when this step works? | How bad is it when this step is wrong? |
"""


def test_an_unnamed_first_column_survives_the_round_trip(tmp_path):
    """Two of the 53 masters lead with an unnamed column. Reading headers as "cells that have a
    value" dropped it, shifted every row one to the left and silently lost the last column — the
    kind of corruption that still looks like a table."""
    parsed = master.parse(UNNAMED_FIRST)
    assert parsed.headers[0] == "", "the fixture must really have an unnamed first column"
    path = tmp_path / "a.xlsx"
    wb.export({"risk-derivation": (parsed, ("property", "Exposure", "risk"))}, path)
    back = wb.read_sheet(path, "risk-derivation")
    assert back.headers == parsed.headers
    assert back.rows == parsed.rows


def test_the_master_s_whole_meta_block_survives_not_just_its_name(tmp_path):
    """`Section` is what keys the derived JSON and `Source` is the artifact's provenance. Carrying
    only `Artifact` would have published records under the wrong key and lost where they came
    from."""
    parsed = master.parse(UNNAMED_FIRST)
    path = tmp_path / "a.xlsx"
    wb.export({"risk-derivation": (parsed, ("property", "Exposure", "risk"))}, path)
    back = wb.read_sheet(path, "risk-derivation")
    assert back.meta.get("Section") == "properties"
    assert back.meta.get("Source") == "CAFE_Artifacts_Visualisation_v0_25.html"
    assert "Record type" not in back.meta, "the catalogue rows are ours, not the master's"


# ------------------------------------------------- how each artifact is meant to be READ

def test_the_retrieval_mode_is_computed_the_way_the_publisher_computes_it():
    """`whole`, `key` and `vector` are not cosmetic: they change what good content IS. A `whole`
    register is read in full, so every row must earn its place and "the relevant rows" would
    silently drop a rule. A `key` artifact is read by exact natural key, so that key must be
    unique and stable. A `vector` artifact is searched for relevance, so its text has to be worth
    embedding — prose, not a code.

    Computed rather than restated: declared modes come from the publisher's own RETRIEVAL table,
    and an artifact that declares nothing takes the kind's default (prose can only be searched; a
    record keeps its exact read). A second hand-maintained copy would drift from the thing that
    actually publishes."""
    assert wb.retrieval_for("criticality-taxonomy", "class") == "whole"     # declared
    assert wb.retrieval_for("ai-capability-map", "capability") == "whole"   # declared
    assert wb.retrieval_for("guardrails", "guardrail") == "key"             # record default
    assert wb.retrieval_for("some-prose-artifact", "") == "vector"          # prose default


def test_every_mode_carries_the_guidance_that_makes_it_actionable():
    """A mode a reader cannot act on is a label. The team is being asked to WRITE these, so each
    mode says what it means for the content."""
    for mode in ("whole", "key", "vector"):
        assert wb.RETRIEVAL_NOTE[mode] and len(wb.RETRIEVAL_NOTE[mode]) > 40


def test_the_sheet_and_the_index_both_say_how_the_artifact_is_read(tmp_path):
    parsed = master.parse(SAMPLE)
    path = tmp_path / "a.xlsx"
    wb.export({"criticality-taxonomy": (parsed, ("class", "Class", "risk"))}, path)
    import openpyxl
    book = openpyxl.load_workbook(path)
    sheet_cells = [str(c.value) for row in book["criticality-taxonomy"].iter_rows() for c in row]
    assert "whole" in sheet_cells
    assert any(wb.RETRIEVAL_NOTE["whole"][:30] in str(c) for c in sheet_cells), "guidance too"
    index_cells = [str(c.value) for row in book[wb.INDEX].iter_rows() for c in row]
    assert "whole" in index_cells


def test_the_read_mode_is_ours_and_never_becomes_part_of_the_master(tmp_path):
    """It is catalogue metadata, not artifact content. Writing it into the master's meta block
    would publish a field the publisher did not put there."""
    parsed = master.parse(SAMPLE)
    path = tmp_path / "a.xlsx"
    wb.export({"criticality-taxonomy": (parsed, ("class", "Class", "risk"))}, path)
    back = wb.read_sheet(path, "criticality-taxonomy")
    assert "Read as" not in back.meta and back.meta == parsed.meta


def test_the_index_shows_the_vector_artifacts_that_have_no_sheet(tmp_path):
    """All 53 masters are RECORD artifacts — `whole` or `key`. The vector-mode artifacts are the
    licensed capability workbooks, which are never files in this repository and so can have no
    sheet. Leaving them out entirely would tell the team the corpus has no searchable artifacts,
    which is false and is exactly the kind of silence this codebase keeps paying for. They appear
    as rows with no sheet and a reason."""
    path = tmp_path / "a.xlsx"
    wb.export({"criticality-taxonomy": (master.parse(SAMPLE), ("class", "Class", "risk"))}, path)
    import openpyxl
    rows = list(openpyxl.load_workbook(path)[wb.INDEX].iter_rows(min_row=2, values_only=True))
    vector = [r for r in rows if r[1] == "vector"]
    assert vector, "no vector artifact is listed at all"
    assert all(not r[-1] for r in vector), "a licensed workbook has no sheet to point at"
    assert any("licen" in str(c).lower() for r in vector for c in r), "and it says why"
