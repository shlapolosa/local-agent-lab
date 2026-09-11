"""A capability workbook as a master — on a SYNTHETIC workbook, never a licensed one.

The property that matters is the KEY: `(id, parent, level)` is what lets an exact read ask for a
node's children or the top level by JSONB containment, so every row's parent must exist (or be the
root sentinel) and its level must be its depth in the path.
"""
import io

import openpyxl
import pytest

from lab.core.reference.workbook import HEADERS, KEY_FIELDS, ROOT, TEXT_FIELDS, capability_table


def workbook(rows=None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Capability Map"
    ws.append(["Capability Map"])
    ws.append(["Tier", "Level", "Capability", "Definition"])
    for row in rows or [[1, 1, "Care Delivery", "Delivering care to patients"],
                        ["", 2, "Triage", "Sorting patients by urgency"],
                        ["", 3, "Urgent triage", "The urgent path"],
                        [2, 1, "Finance", "Money"],
                        ["", 2, "Billing", "Invoicing payers"]]:
        ws.append(row)
    other = wb.create_sheet("Value Stream Inventory")
    other.append(["Value Stream Name", "Definition"])
    other.append(["Acquire patient", "not a capability"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_the_table_has_the_declared_columns_and_only_capabilities():
    master = capability_table(workbook(), scheme="test-scheme", title="Test scheme")
    assert master.headers == HEADERS and master.is_table
    assert [r[HEADERS.index("label")] for r in master.rows] == [
        "Care Delivery", "Triage", "Urgent triage", "Finance", "Billing"]
    assert set(KEY_FIELDS) <= set(HEADERS) and set(TEXT_FIELDS) <= set(HEADERS)


def test_every_parent_exists_or_is_the_root_and_level_is_the_depth_of_the_path():
    master = capability_table(workbook(), scheme="test-scheme", title="Test scheme")
    by = {r[0]: dict(zip(HEADERS, r)) for r in master.rows}
    for row in by.values():
        assert row["parent"] == ROOT or row["parent"] in by, row
        assert int(row["level"]) == len(row["path"].split(" > ")), row
    triage = next(r for r in by.values() if r["label"] == "Urgent triage")
    assert triage["path"] == "Care Delivery > Triage > Urgent triage"
    assert by[triage["parent"]]["label"] == "Triage"


def test_ids_are_the_semantic_layer_s_own_so_a_corpus_row_and_a_scheme_concept_agree():
    from lab.core.semantic.skos import concept_id
    master = capability_table(workbook(), scheme="test-scheme", title="Test scheme")
    ids = {r[0] for r in master.rows}
    assert concept_id("test-scheme", ("Care Delivery", "Triage")) in ids
    assert concept_id("other-scheme", ("Care Delivery", "Triage")) not in ids


def test_the_tier_is_kept_where_the_workbook_gives_one():
    master = capability_table(workbook(), scheme="s", title="S")
    tiers = {dict(zip(HEADERS, r))["label"]: dict(zip(HEADERS, r))["tier"] for r in master.rows}
    assert tiers["Care Delivery"] == "1" and tiers["Triage"] == ""


def test_each_row_carries_the_context_a_person_uses_to_place_it():
    """A passage of "path + definition" embeds a leaf on its own words. A person placing a
    capability also reads what its parent MEANS and what sits next to it; the context column
    gives the embedding the same — the parent's first sentence and the sibling labels."""
    table = capability_table(workbook([[1, 1, "Care Delivery", "Delivering care to patients. Also more."],
                                       ["", 2, "Triage", "Sorting patients by urgency"],
                                       ["", 3, "Urgent triage", "The urgent path"],
                                       ["", 2, "Admission", "Taking a patient in"]]),
                             scheme="s", title="t")
    rows = {r[HEADERS.index("label")]: dict(zip(HEADERS, r)) for r in table.rows}
    assert "context" in HEADERS and "context" in TEXT_FIELDS
    assert rows["Triage"]["context"] == "under Care Delivery: Delivering care to patients | beside: Admission"
    assert rows["Urgent triage"]["context"] == "under Triage: Sorting patients by urgency"
    assert rows["Care Delivery"]["context"] == "", "a top-level capability has no parent to cite"


def test_a_workbook_with_no_capability_map_derives_nothing_rather_than_something_else():
    wb = openpyxl.Workbook()
    wb.active.title = "Notes"
    buf = io.BytesIO()
    wb.save(buf)
    master = capability_table(buf.getvalue(), scheme="s", title="S")
    assert master.rows == ()
